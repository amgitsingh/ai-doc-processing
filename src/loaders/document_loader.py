import os
from dataclasses import dataclass
from typing import List

import pypdf
import docx

from src.utils.logger import get_logger

logger = get_logger(__name__)


@dataclass
class LoadedDocument:
    filename: str
    filepath: str
    text: str


def load_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def load_pdf(path: str) -> str:
    reader = pypdf.PdfReader(path)
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def load_docx(path: str) -> str:
    document = docx.Document(path)
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


LOADERS = {
    ".txt": load_txt,
    ".pdf": load_pdf,
    ".docx": load_docx,
}


def load_document(path: str) -> LoadedDocument:
    """
    Load a single document file into plain text.
    """
    if not os.path.exists(path):
        raise FileNotFoundError(f"File not found: {path}")

    ext = os.path.splitext(path)[1].lower()
    loader_fn = LOADERS.get(ext)

    if loader_fn is None:
        raise ValueError(f"Unsupported file type: {ext}")

    text = loader_fn(path)

    if not text.strip():
        raise ValueError(f"No extractable text found in file: {path}")

    return LoadedDocument(
        filename=os.path.basename(path),
        filepath=os.path.abspath(path),
        text=text,
    )


def load_documents(folder: str) -> List[LoadedDocument]:
    """
    Load every supported document in `folder` for batch processing.
    Any file that fails to load (unsupported type, corrupted file,
    empty extraction, etc.) is logged and skipped rather than crashing
    the whole batch.
    """
    if not os.path.isdir(folder):
        raise FileNotFoundError(f"Data folder not found: {folder}")

    loaded: List[LoadedDocument] = []

    for filename in sorted(os.listdir(folder)):
        path = os.path.join(folder, filename)

        if not os.path.isfile(path):
            continue

        try:
            document = load_document(path)
            loaded.append(document)
            logger.info(f"Loaded document: {filename} ({len(document.text)} chars)")
        except Exception as e:
            logger.error(f"Failed to load '{filename}': {e}")
            continue

    logger.info(f"Loaded {len(loaded)} of {len(os.listdir(folder))} file(s) from {folder}")
    return loaded
